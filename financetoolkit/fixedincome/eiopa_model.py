"""European Insurance and Occupational Pensions Authority (EIOPA) Model"""

__docformat__ = "google"

import io
import re
import zipfile

import numpy as np
import openpyxl
import pandas as pd

from financetoolkit import helpers
from financetoolkit.cache import policy_model
from financetoolkit.economics.helpers import collect_cached_data
from financetoolkit.utilities.excel_model import read_excel
from financetoolkit.utilities.logger_model import get_logger
from financetoolkit.utilities.requests_model import get_request

logger = get_logger()

BASE_URL = "https://www.eiopa.europa.eu"

# The page that links every monthly release published since January 2023. A release is a
# zip file whose link carries a random identifier, so the links are read from this page.
RELEASES_PAGE = f"{BASE_URL}/tools-and-data/risk-free-interest-rate-term-structures_en"
RELEASE_PATTERN = re.compile(r'href="([^"]*?filename=EIOPA_RFR_(\d{8})\.zip)"')

# The worksheets of the term structures workbook, per curve.
CURVES = {
    "spot_no_va": "RFR_spot_no_VA",
    "spot_with_va": "RFR_spot_with_VA",
}

# The shocked curves of the standard formula's interest rate risk submodule. EIOPA ships
# their worksheets as formulas without computed values (every cell reads 0.01 until Excel
# recalculates), so they are computed here from the curve without volatility adjustment
# and the shocks, with the formulas of those worksheets.
SHOCKED_CURVES = {"shock_up": "up", "shock_down": "down"}
SHOCKS_SHEET = "Shocks"

# The columns of the Shocks worksheet with the maturity and the downward and upward shock,
# and the first row with a maturity. Only some maturities hold a value; the others are
# formulas that interpolate linearly between them and stay flat after the last.
SHOCK_COLUMNS = {"maturity": 1, "down": 3, "up": 4}
FIRST_SHOCK_ROW = 10

# The shocked rate is rounded like EIOPA's worksheets, and an upward shock is at least
# one percentage point.
SHOCK_DECIMALS = 5
MINIMUM_UPWARD_SHOCK = 0.01

# The names EIOPA gives the currencies and countries that differ from the ones used
# elsewhere in the Finance Toolkit.
NAMES = {"Euro": "Euro Area", "Czechia": "Czech Republic"}

# The row of the worksheet with the names, and the first row with a maturity.
NAME_ROW = 1
FIRST_MATURITY_ROW = 10


def get_releases() -> pd.DataFrame:
    """
    Lists the monthly releases EIOPA publishes, with the link to each.

    Returns:
        pd.DataFrame: The link per release, indexed by the month the release is for.

    Raises:
        ValueError: When the page links no releases, which means it changed.
    """
    description = "EIOPA risk-free rate releases"

    def fetch() -> pd.DataFrame:
        page = get_request(RELEASES_PAGE, timeout=60).text
        releases = {
            pd.Period(pd.Timestamp(date), freq="M"): link
            for link, date in RELEASE_PATTERN.findall(page)
        }

        if not releases:
            raise ValueError(
                f"The {description} page links no releases, which means EIOPA changed it."
            )

        frame = pd.DataFrame({"Link": list(releases.values())}, index=list(releases))
        frame.index = pd.PeriodIndex(frame.index, freq="M")

        return frame.sort_index()

    return collect_cached_data(
        source=policy_model.EIOPA,
        dataset="releases",
        entity="releases",
        fetch=fetch,
        description=description,
    )


def _parse_term_structures(
    workbook: bytes, sheet: str, description: str
) -> pd.DataFrame:
    """
    Reads one curve from the term structures workbook of a release.

    Args:
        workbook (bytes): The workbook.
        sheet (str): The worksheet of the curve, e.g. "RFR_spot_no_VA".
        description (str): What is read, used in the error message.

    Returns:
        pd.DataFrame: The rates as decimals, indexed by maturity in years with a column per
        currency or country.

    Raises:
        ValueError: When the worksheet is missing or has no maturities where expected.
    """
    try:
        data = read_excel(workbook, sheet_name=sheet, header=None)
    except ValueError as error:
        raise ValueError(
            f"The {description} has no worksheet '{sheet}', which means EIOPA changed its layout."
        ) from error

    names = data.iloc[NAME_ROW, 2:]
    rates = data.iloc[FIRST_MATURITY_ROW:, 1:]
    rates = rates[pd.to_numeric(rates.iloc[:, 0], errors="coerce").notna()]

    if rates.empty:
        raise ValueError(
            f"The {description} has no maturities where they are expected, which means EIOPA "
            "changed its layout."
        )

    curve = rates.iloc[:, 1:].apply(pd.to_numeric, errors="coerce")
    curve.index = rates.iloc[:, 0].astype(int).to_numpy()
    curve.columns = [NAMES.get(str(name).strip(), str(name).strip()) for name in names]

    return curve.loc[:, [column for column in curve.columns if column != "nan"]]


def _parse_shocks(workbook: bytes, description: str) -> pd.DataFrame:
    """
    Reads the relative downward and upward shocks per maturity from the Shocks worksheet.
    The worksheet holds values for some maturities and formulas for the others, which
    interpolate linearly between them and stay flat after the last, so only the values are
    read and the formulas are evaluated here.

    Args:
        workbook (bytes): The workbook.
        description (str): What is read, used in the error message.

    Returns:
        pd.DataFrame: The "down" and "up" shocks, indexed by maturity in years.

    Raises:
        ValueError: When the worksheet is missing or holds no shocks.
    """
    book = openpyxl.load_workbook(io.BytesIO(workbook), read_only=True, data_only=False)

    if SHOCKS_SHEET not in book.sheetnames:
        raise ValueError(
            f"The {description} has no worksheet '{SHOCKS_SHEET}', which means EIOPA changed "
            "its layout."
        )

    maturities, values = [], {"down": [], "up": []}

    for row in book[SHOCKS_SHEET].iter_rows(
        min_row=FIRST_SHOCK_ROW + 1, values_only=True
    ):
        maturity = row[SHOCK_COLUMNS["maturity"]]

        if not isinstance(maturity, int | float):
            continue

        maturities.append(int(maturity))
        for direction, shocks in values.items():
            value = row[SHOCK_COLUMNS[direction]]
            # A formula is read as its text, so only numbers are values.
            shocks.append(float(value) if isinstance(value, int | float) else np.nan)

    book.close()
    shocks = pd.DataFrame(values, index=maturities)

    if shocks.dropna(how="all").empty:
        raise ValueError(
            f"The {description} has no interest rate shocks, which means EIOPA changed its layout."
        )

    return shocks.interpolate(method="index", limit_area="inside").ffill()


def _apply_shock(
    curve: pd.DataFrame, shocks: pd.DataFrame, direction: str
) -> pd.DataFrame:
    """
    Applies the standard formula's interest rate shock to a curve, as EIOPA's worksheets
    do: upwards by the relative shock and at least one percentage point, downwards by the
    relative shock, leaving negative rates unchanged.

    Args:
        curve (pd.DataFrame): The rates, indexed by maturity in years.
        shocks (pd.DataFrame): The "down" and "up" shocks, indexed by maturity in years.
        direction (str): "up" or "down".

    Returns:
        pd.DataFrame: The shocked rates.
    """
    shock = shocks[direction].reindex(curve.index).to_numpy()[:, None]
    rates = curve.to_numpy(dtype=float)

    if direction == "up":
        shocked = rates + np.maximum(MINIMUM_UPWARD_SHOCK, shock * np.abs(rates))
    else:
        shocked = np.where(rates < 0, rates, rates - shock * np.abs(rates))

    return pd.DataFrame(
        np.round(shocked, SHOCK_DECIMALS), index=curve.index, columns=curve.columns
    )


def _get_release(month: pd.Period, link: str, curve: str) -> pd.DataFrame:
    """
    Retrieves one curve of one monthly release. A published release does not change, so
    it is cached for a year.

    Args:
        month (pd.Period): The month of the release.
        link (str): The link to the release's zip file.
        curve (str): The curve, a key of CURVES or SHOCKED_CURVES.

    Returns:
        pd.DataFrame: The rates, indexed by maturity in years with a column per currency or
        country.
    """
    description = f"EIOPA risk-free rates of {month}"

    def fetch() -> pd.DataFrame:
        archive = zipfile.ZipFile(
            io.BytesIO(get_request(f"{BASE_URL}{link}", timeout=180).content)
        )
        workbook = next(
            (
                name
                for name in archive.namelist()
                if name.endswith("_Term_Structures.xlsx")
            ),
            None,
        )

        if workbook is None:
            raise ValueError(
                f"The {description} have no term structures workbook, which means EIOPA "
                "changed the release."
            )

        content = archive.read(workbook)

        if curve not in SHOCKED_CURVES:
            return _parse_term_structures(content, CURVES[curve], description)

        return _apply_shock(
            _parse_term_structures(content, CURVES["spot_no_va"], description),
            _parse_shocks(content, description),
            SHOCKED_CURVES[curve],
        )

    return collect_cached_data(
        source=policy_model.EIOPA,
        dataset="release",
        # The shocked curves were read from their worksheets before, which gave nothing.
        entity=f"{month}/{curve}" + ("/computed" if curve in SHOCKED_CURVES else ""),
        fetch=fetch,
        description=description,
    )


def get_risk_free_rate_term_structures(
    curve: str, start_date: str, end_date: str
) -> pd.DataFrame:
    """
    Retrieves the risk-free rate term structures EIOPA publishes monthly for Solvency II,
    for every currency and country it covers and maturities of 1 to 150 years. Only the
    releases between the start and end date are downloaded.

    Args:
        curve (str): "spot_no_va", "spot_with_va", "shock_up" or "shock_down".
        start_date (str): The start date (YYYY-MM-DD).
        end_date (str): The end date (YYYY-MM-DD).

    Returns:
        pd.DataFrame: The rates as decimals, indexed by the month of the release with a
        column per currency or country and maturity.
    """
    releases = get_releases()

    if releases.empty:
        return pd.DataFrame()

    selected = releases.loc[pd.Period(start_date, "M") : pd.Period(end_date, "M")]

    if selected.empty:
        logger.warning(
            "EIOPA publishes its risk-free rates monthly from %s; there is no release between "
            "%s and %s.",
            releases.index[0],
            start_date,
            end_date,
        )
        return pd.DataFrame()

    curves = helpers.run_in_parallel(
        _get_release,
        [(month, link, curve) for month, link in selected["Link"].items()],
        max_workers=4,
    )

    frames = {
        month: frame.stack()
        for month, frame in zip(selected.index, curves, strict=True)
        if not frame.empty
    }

    if not frames:
        return pd.DataFrame()

    term_structures = pd.DataFrame(frames).T
    term_structures.columns = term_structures.columns.swaplevel(0, 1)
    term_structures = term_structures.sort_index(
        axis=1, level=[0, 1], sort_remaining=False
    )

    # Maturities in years, labelled like the other curves of the Finance Toolkit.
    term_structures.columns = pd.MultiIndex.from_tuples(
        [(country, f"{maturity}Y") for country, maturity in term_structures.columns],
        names=["Country", "Maturity"],
    )
    term_structures.index = pd.PeriodIndex(term_structures.index, freq="M")

    return term_structures


# The symmetric adjustment of the equity capital charge (the "equity dampener") EIOPA
# publishes monthly: each workbook holds the full daily history from 1991 on its
# Calculations worksheet, so only the latest is read.
SYMMETRIC_ADJUSTMENT_PAGE = (
    f"{BASE_URL}/tools-and-data/symmetric-adjustment-equity-capital-charge_en"
)
SYMMETRIC_ADJUSTMENT_PATTERN = re.compile(
    r'href="([^"]*filename=EIOPA(?:_|%20)symmetric(?:_|%20)adjustment[^"]*\.xlsx)"',
    flags=re.IGNORECASE,
)
SYMMETRIC_ADJUSTMENT_SHEET = "Calculations"
SYMMETRIC_ADJUSTMENT_DATE = "All calendar days"
SYMMETRIC_ADJUSTMENT_VALUE = "Dampener final"

# The standard equity charges of the Solvency II standard formula before the symmetric
# adjustment: type 1 (listed in the EEA or OECD) and type 2 (other) equities.
EQUITY_CHARGES = {"Type 1 Equity Charge": 0.39, "Type 2 Equity Charge": 0.49}


def get_symmetric_adjustment() -> pd.DataFrame:
    """
    Retrieves the daily symmetric adjustment of the equity capital charge of Solvency II,
    from 1991, and the type 1 and type 2 equity charges it results in.

    Returns:
        pd.DataFrame: The "Symmetric Adjustment" and the equity charges as decimals, indexed
        by day.

    Raises:
        ValueError: When the page links no workbook or the workbook misses the history.
    """
    description = "EIOPA symmetric adjustment of the equity capital charge"

    def fetch() -> pd.DataFrame:
        page = get_request(SYMMETRIC_ADJUSTMENT_PAGE, timeout=60).text
        links = SYMMETRIC_ADJUSTMENT_PATTERN.findall(page)

        if not links:
            raise ValueError(
                f"The {description} page links no workbook, which means EIOPA changed it."
            )

        # The page lists the latest month first.
        link = links[0] if links[0].startswith("http") else f"{BASE_URL}{links[0]}"

        sheet = read_excel(
            get_request(link, timeout=120).content,
            sheet_name=SYMMETRIC_ADJUSTMENT_SHEET,
            header=None,
        )

        header = sheet.index[
            sheet.astype(str).eq(SYMMETRIC_ADJUSTMENT_DATE).any(axis=1)
        ]
        if header.empty:
            raise ValueError(
                f"The {description} misses its daily history, which means EIOPA changed the "
                "workbook."
            )

        names = sheet.loc[header[0]].astype(str)
        date_column = names[names == SYMMETRIC_ADJUSTMENT_DATE].index[0]
        value_column = names[names == SYMMETRIC_ADJUSTMENT_VALUE].index[0]
        data = sheet.loc[header[0] + 1 :, [date_column, value_column]].dropna()
        data.columns = ["Date", "Symmetric Adjustment"]
        data["Date"] = pd.to_datetime(data["Date"], errors="coerce")
        data = data.dropna(subset=["Date"])

        adjustment = pd.DataFrame(
            {
                "Symmetric Adjustment": pd.to_numeric(
                    data["Symmetric Adjustment"], errors="coerce"
                ).to_numpy()
            },
            index=pd.PeriodIndex(data["Date"], freq="D"),
        ).sort_index()

        for charge, base in EQUITY_CHARGES.items():
            adjustment[charge] = base + adjustment["Symmetric Adjustment"]

        return adjustment

    return collect_cached_data(
        source=policy_model.EIOPA,
        dataset="symmetric_adjustment",
        entity="latest",
        fetch=fetch,
        description=description,
    )
