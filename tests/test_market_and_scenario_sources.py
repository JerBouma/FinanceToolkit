"""Offline tests for the volatility, stress test, corporate bond and carbon price sources."""

import io

import numpy as np
import pandas as pd
import pytest
import requests

from financetoolkit.cache.cache_controller import clear_active_cache


@pytest.fixture(autouse=True)
def isolate_from_the_active_cache():
    """Start every test without an active cache another test may have left behind."""
    clear_active_cache()
    yield
    clear_active_cache()


class FakeResponse:
    def __init__(self, text="", content=None, status_code=200):
        self.text = text
        self.content = content if content is not None else text.encode()
        self.status_code = status_code
        self.headers = {}


def _workbook(sheets: dict[str, list[list]]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(
                writer, sheet_name=name, header=False, index=False
            )
    return buffer.getvalue()


def _not_found(url):
    response = requests.Response()
    response.status_code = 404
    return requests.exceptions.HTTPError(f"404 for {url}", response=response)


def test_implied_volatility_is_a_decimal_per_market_and_maturity(monkeypatch):
    from financetoolkit import Economics
    from financetoolkit.economics import cboe_model, fmp_model, stoxx_model

    days = ("2026-09-29", "2026-09-30")
    cboe = {
        "VIX9D": "DATE,OPEN,HIGH,LOW,CLOSE\n09/29/2026,1,1,1,13.0\n09/30/2026,1,1,1,14.2\n",
        "VIX": "DATE,OPEN,HIGH,LOW,CLOSE\n09/29/2026,1,1,1,15.5\n09/30/2026,1,1,1,16.3\n",
        "VVIX": "DATE,VVIX\n09/29/2026,85.4\n09/30/2026,86.0\n",
    }

    def fake_cboe(url, timeout):
        symbol = url.rsplit("/", 1)[-1].removesuffix("_History.csv")
        if symbol not in cboe:
            raise _not_found(url)
        return FakeResponse(text=cboe[symbol])

    vstoxx = "Date;Symbol;Indexvalue\n29.09.2026;V6I1;17.6\n30.09.2026;V6I1;18.1\n"
    monkeypatch.setattr(cboe_model, "get_request", fake_cboe)
    monkeypatch.setattr(
        stoxx_model,
        "get_request",
        lambda url, timeout: FakeResponse(
            text=vstoxx if url.endswith("v6i1.txt") else ""
        ),
    )
    monkeypatch.setattr(
        fmp_model,
        "get_index_history",
        lambda symbol, api_key, start, end: pd.Series(
            [105.2, 110.4] if symbol == "^MOVE" else [4.5, 4.6],
            index=pd.PeriodIndex(list(days), freq="D"),
        ),
    )

    economics = Economics(start_date="2026-09-01", end_date="2026-09-30", api_key="key")
    volatility = economics.get_implied_volatility(
        markets=[
            "US Equity",
            "Euro Area Equity",
            "US Treasury Rates",
            "US Equity Volatility",
        ],
        rounding=6,
    )
    last = volatility.loc[pd.Period("2026-09-30", "D")]

    assert last[("US Equity", "9D")] == pytest.approx(0.142)
    assert last[("US Equity", "1M")] == pytest.approx(0.163)
    assert last[("Euro Area Equity", "1M")] == pytest.approx(0.181)
    assert last[("US Equity Volatility", "1M")] == pytest.approx(0.86)
    # The MOVE is in basis points: 110.4 basis points is 0.01104.
    assert last[("US Treasury Rates", "1M")] == pytest.approx(0.01104)
    assert volatility.columns.names == ["Market", "Maturity"]

    # Without an API key the FMP markets are left out.
    without_key = Economics(start_date="2026-09-01", end_date="2026-09-30")
    assert without_key.get_implied_volatility(markets="US Treasury Rates").empty


def test_federal_reserve_scenario_is_split_by_region_in_decimals(monkeypatch):
    from financetoolkit.economics import frb_model

    domestic = (
        "Scenario Name,Date,Real GDP growth,10-year Treasury yield,"
        "House Price Index (Level),Market Volatility Index (Level)\n"
        "Supervisory Severely Adverse,2026 Q1,-5.4,3.1,303,59.7\n"
        "Supervisory Severely Adverse,2026 Q2,-4.9,2.7,283.2,72\n"
    )
    international = (
        "Scenario Name,Date,Euro area real GDP growth,"
        "Euro area bilateral dollar exchange rate (USD/euro),U.K. inflation\n"
        "Supervisory Severely Adverse,2026 Q1,-8.6,1.124,0.9\n"
        "Supervisory Severely Adverse,2026 Q2,-8.5,1.08,-0.1\n"
    )
    requested = []

    def fake(url, timeout):
        requested.append(url)
        if "2026_Final" not in url:
            raise _not_found(url)
        return FakeResponse(
            text=domestic if url.endswith("Domestic.csv") else international
        )

    monkeypatch.setattr(frb_model, "get_request", fake)

    scenario = frb_model.get_supervisory_scenario("adverse", year=2026)
    first = scenario.loc[pd.Period("2026Q1", "Q")]

    assert first[("United States", "Real GDP growth")] == pytest.approx(-0.054)
    assert first[("United States", "House Price Index (Level)")] == pytest.approx(303)
    assert first[("United States", "Market Volatility Index (VIX)")] == pytest.approx(
        0.597
    )
    assert first[("Euro Area", "Real GDP growth")] == pytest.approx(-0.086)
    assert first[
        ("Euro Area", "Bilateral dollar exchange rate (USD/euro)")
    ] == pytest.approx(1.124)
    assert first[("United Kingdom", "Inflation")] == pytest.approx(0.009)
    # A year whose scenarios are not published gives nothing rather than an error.
    assert frb_model.get_supervisory_scenario("adverse", year=2031).empty


def test_esrb_scenario_reads_the_country_panel(monkeypatch):
    from financetoolkit.economics import esrb_model

    gdp = [[None] * 9 for _ in range(3)]
    gdp[1] = [None, "Real GDP"] + [None] * 7
    gdp += [
        [
            None,
            None,
            None,
            "Historical growth (%)",
            "Baseline growth (%)",
            None,
            "Adverse growth (%)",
            None,
            None,
        ],
        [None, None, None, 2024, 2025, 2026, 2025, 2026, None],
        [None, "Germany", "DE", -0.16, 0.16, 0.84, -3.63, -4.24, None],
        [None, "Euro area", "EA", 0.8, 0.9, 1.2, -2.8, -3.5, None],
        [None, "Notes: annual averages.", None, None, None, None, None, None, None],
    ]
    itraxx = [[None] * 9 for _ in range(3)]
    itraxx += [
        [
            None,
            None,
            "Historical level",
            "Baseline level",
            None,
            "Adverse level",
            None,
            None,
            None,
        ],
        [None, None, 2024, 2025, 2026, 2025, 2026, None, None],
        [None, "iTraxx Overall 5y", 60, 62, 64, 240, 198, None, None],
    ]
    workbook = _workbook({"GDP": gdp, "Itraxx": itraxx})
    page = (
        '<a href="/mppa/stress/shared/pdf/esrb.stress_test230131.macrofinancialscenario~a.en.xlsx">2023</a>'
        '<a href="/mppa/stress/shared/pdf/esrb.stress_test250120.macrofinancialscenario~b.en.xlsx">2025</a>'
    )
    requested = []

    def fake(url, timeout):
        requested.append(url)
        return (
            FakeResponse(text=page)
            if url.endswith(".html")
            else FakeResponse(content=workbook)
        )

    monkeypatch.setattr(esrb_model, "get_request", fake)

    adverse = esrb_model.get_macro_financial_scenario("adverse")

    assert requested[-1].endswith("stress_test250120.macrofinancialscenario~b.en.xlsx")
    assert adverse.loc[
        pd.Period("2026", "Y"), ("Germany", "Real GDP growth")
    ] == pytest.approx(-0.0424)
    assert adverse.loc[
        pd.Period("2025", "Y"), ("Euro Area", "Real GDP growth")
    ] == pytest.approx(-0.028)
    # iTraxx spreads are in basis points.
    assert adverse.loc[
        pd.Period("2025", "Y"), ("iTraxx Overall 5y", "Credit spread")
    ] == pytest.approx(0.024)
    assert not any(
        label.startswith("Notes") for label in adverse.columns.get_level_values(0)
    )

    historic = esrb_model.get_macro_financial_scenario("historic")
    assert list(historic.index) == [pd.Period("2024", "Y")]


def test_corporate_bond_yields_and_spreads(monkeypatch):
    from financetoolkit import FixedIncome
    from financetoolkit.fixedincome import bundesbank_model, rba_model

    bundesbank = (
        "BBK_SEIS_ISSUER_CLASS;TIME_PERIOD;OBS_VALUE;OBS_STATUS\n"
        "X2000;2026-07;3.96;\nX2000;2026-08;4.06;\n"
        "S13;2026-07;3.11;\nS13;2026-08;3.21;\nS13;1955-01;0.00;N\n"
    )
    monkeypatch.setattr(
        bundesbank_model,
        "get_request",
        lambda url, timeout, extra_headers: FakeResponse(text=bundesbank),
    )
    f3 = (
        "F3 AGGREGATE MEASURES\nTitle,a,b\nSeries ID,FNFYA5M,FNFYBBB7M\n"
        "31/07/2026,5.41,5.70\n31/08/2026,5.55,5.82\n"
    )
    f2 = (
        "F2 CAPITAL MARKET YIELDS\nSeries ID,FCMYGBAG2D,FCMYGBAG3D,FCMYGBAG5D,FCMYGBAG10D\n"
        "31-Jul-2026,4.5,4.5,4.6,4.9\n31-Aug-2026,4.6,4.6,4.7,5.1\n"
    )
    monkeypatch.setattr(
        rba_model,
        "get_request",
        lambda url, timeout, extra_headers: FakeResponse(
            text=f3 if "f3-" in url else f2
        ),
    )

    fixedincome = FixedIncome(start_date="2026-07-01", end_date="2026-08-31")
    yields = fixedincome.get_corporate_bond_yields()
    spreads = fixedincome.get_corporate_bond_yields(spread=True)
    august = pd.Period("2026-08", "M")

    assert yields.loc[august, ("Germany", "All ratings")] == pytest.approx(0.0406)
    assert yields.loc[august, ("Australia", "BBB 7Y")] == pytest.approx(0.0582)
    assert spreads.loc[august, ("Germany", "All ratings")] == pytest.approx(0.0085)
    assert spreads.loc[august, ("Australia", "A 5Y")] == pytest.approx(0.0085)
    # 7 years lies between the 5 and 10-year government yields: 4.7% + 2/5 of 0.4%.
    assert spreads.loc[august, ("Australia", "BBB 7Y")] == pytest.approx(
        0.0582 - 0.0486
    )
    assert "United States" not in yields.columns.get_level_values(0)


def test_carbon_price_averages_the_days_general_allowance_auctions(monkeypatch):
    from financetoolkit import Economics
    from financetoolkit.economics import eex_model

    rows = [[None] * 7 for _ in range(5)]
    rows += [
        [
            None,
            "Date",
            "Time",
            "Auction Name",
            "Contract",
            "Status",
            "Auction Price €/tCO2",
        ],
        [
            None,
            pd.Timestamp("2026-10-06"),
            None,
            "Auction CAP3 EU",
            "T3PA",
            "successful",
            84.31,
        ],
        [
            None,
            pd.Timestamp("2026-10-06"),
            None,
            "Auction DE",
            "T3PA",
            "successful",
            84.11,
        ],
        [
            None,
            pd.Timestamp("2026-10-05"),
            None,
            "Auction CAP3 EU",
            "T3PA",
            "successful",
            82.6,
        ],
        [
            None,
            pd.Timestamp("2026-10-05"),
            None,
            "EUAA Auction",
            "EAA3",
            "successful",
            70.0,
        ],
        [
            None,
            pd.Timestamp("2026-10-02"),
            None,
            "Auction CAP3 EU",
            "T3PA",
            "cancelled",
            None,
        ],
    ]
    workbook = _workbook({"Primary Market Auction": rows})
    requested = []

    def fake(url, timeout):
        requested.append(url)
        if "2026" not in url:
            raise _not_found(url)
        return FakeResponse(content=workbook)

    monkeypatch.setattr(eex_model, "get_request", fake)

    prices = Economics(
        start_date="2026-10-01", end_date="2026-10-06"
    ).get_carbon_price()

    assert prices["European Union"].to_dict() == {
        pd.Period("2026-10-05", "D"): pytest.approx(82.6),
        pd.Period("2026-10-06", "D"): pytest.approx(84.21),
    }
    # Only the years from the buffered start date to the end date are requested.
    assert all("2025" in url or "2026" in url for url in requested)


def test_rating_transition_matrix_is_withdrawal_adjusted(monkeypatch):
    from financetoolkit import FixedIncome
    from financetoolkit.fixedincome import esma_model

    matrix = (
        "BOP\\EOP;AAA;AA;A;D;Withdrawals\n"
        "AAA;6.0;0.0;0.0;0.0;0.0\n"
        "AA;0.0;217.0;11.0;0.0;1.0\n"
        "A;0.0;16.0;1144.0;0.0;32.0\n"
        "D;0.0;0.0;0.0;0.0;0.0\n"
    )
    bodies = []

    class Response(FakeResponse):
        def raise_for_status(self):
            return None

    def fake_post(url, data, headers, timeout):
        bodies.append((url, data))
        return Response(text=matrix)

    monkeypatch.setattr(esma_model.SESSION, "post", fake_post)

    fixedincome = FixedIncome()
    probabilities = fixedincome.get_rating_transition_matrix(agency="S&P", year=2025)
    counts = fixedincome.get_rating_transition_matrix(
        agency="S&P",
        year=2025,
        probabilities=False,
        include_withdrawals=True,
        rounding=0,
    )

    assert '"STPGB"' in bodies[0][1] and '"1735689600000"' in bodies[0][1]
    assert probabilities.loc["AA", "A"] == pytest.approx(11 / 228, abs=1e-4)
    assert probabilities.sum(axis=1).round(3).eq(1).all()
    # A rating without any outstanding at the start has no row.
    assert "D" not in probabilities.index
    assert counts.loc["A", "Withdrawals"] == 32
    assert fixedincome.get_rating_transition_matrix(agency="Unknown agency").empty


def test_default_rates_start_at_the_first_year_with_a_default(monkeypatch):
    from financetoolkit import FixedIncome
    from financetoolkit.fixedincome import esma_model

    def fake_default_rates(agency, rating_type, start, end, region=None):
        if start.year < 2001:
            return pd.DataFrame(
                {"Defaults": [0, 0], "Default Rate": [0.0, 0.0]}, index=["BB", "CCC"]
            )
        return pd.DataFrame(
            {
                "Defaults": [1, 20],
                "Default Rate": [0.002, 0.25 + start.year % 10 / 100],
            },
            index=["BB", "CCC"],
        )

    monkeypatch.setattr(esma_model, "get_default_rates", fake_default_rates)

    default_rates = FixedIncome(
        start_date="1998-01-01", end_date="2003-12-31"
    ).get_default_rates()

    assert list(default_rates.index) == [
        pd.Period(str(year), "Y") for year in range(2001, 2004)
    ]
    assert default_rates.loc[pd.Period("2003", "Y"), "CCC"] == pytest.approx(0.28)


def test_climate_scenario_applies_the_deviation_to_the_baseline(monkeypatch):
    from financetoolkit import Economics
    from financetoolkit.economics import ngfs_model

    runs = pd.DataFrame(
        {
            "run_id": [1, 2, 3],
            "model": [
                "NiGEM NGFS v1.24.2[REMIND-MAgPIE 3.3-4.8]",
                "NiGEM NGFS v1.24.2[REMIND-MAgPIE 3.3-4.8]",
                "REMIND-MAgPIE 3.3-4.8",
            ],
            "scenario": ["Baseline", "Net Zero 2050", "Net Zero 2050"],
        }
    )
    years = pd.PeriodIndex(["2030", "2050"], freq="Y")
    series = {
        (1, "Long term interest rate ; %"): pd.DataFrame(
            {"Germany": [3.0, 3.2]}, index=years
        ),
        (2, "Long term interest rate ; %(combined)"): pd.DataFrame(
            {"Germany": [0.9, 0.6]}, index=years
        ),
        (1, "Gross Domestic Product (GDP)"): pd.DataFrame(
            {"Germany": [4500.0, 5400.0]}, index=years
        ),
        (2, "Gross Domestic Product (GDP)(combined)"): pd.DataFrame(
            {"Germany": [-2.0, -5.0]}, index=years
        ),
        (3, "Price|Carbon"): pd.DataFrame({"World": [183.3, 748.8]}, index=years),
    }
    monkeypatch.setattr(ngfs_model, "get_runs", lambda: runs)
    monkeypatch.setattr(
        ngfs_model, "_get_timeseries", lambda run, variable: series[(run, variable)]
    )

    economics = Economics()
    rate = economics.get_climate_scenario(countries="Germany")
    rate_deviation = economics.get_climate_scenario(countries="Germany", deviation=True)
    gdp = economics.get_climate_scenario("gdp", countries="Germany")
    carbon = economics.get_climate_scenario("carbon_price")

    assert rate.loc[pd.Period("2030", "Y"), "Germany"] == pytest.approx(0.039)
    assert rate_deviation.loc[pd.Period("2050", "Y"), "Germany"] == pytest.approx(0.006)
    assert gdp.loc[pd.Period("2050", "Y"), "Germany"] == pytest.approx(5400 * 0.95)
    assert carbon.loc[pd.Period("2050", "Y"), "World"] == pytest.approx(748.8)
    assert economics.get_climate_scenario(scenario="Hothouse").empty


def test_long_run_asset_returns_require_accepting_the_licence(monkeypatch):
    from financetoolkit import Economics
    from financetoolkit.economics import macrohistory_model

    data = pd.DataFrame(
        {
            "year": [2019.0, 2020.0, 2019.0, 2020.0],
            "country": ["USA", "USA", "Germany", "Germany"],
            "iso": ["USA", "USA", "DEU", "DEU"],
            "eq_tr": [0.30, 0.20, 0.25, 0.04],
            "cpi": [100.0, 110.0, 100.0, 104.0],
        }
    )
    buffer = io.BytesIO()
    data.to_stata(buffer, write_index=False)
    monkeypatch.setattr(
        macrohistory_model,
        "get_request",
        lambda url, timeout: FakeResponse(content=buffer.getvalue()),
    )

    economics = Economics(start_date="2019-01-01", end_date="2020-12-31")

    assert economics.get_long_run_asset_returns().empty
    nominal = economics.get_long_run_asset_returns(accept_licence=True)
    real = economics.get_long_run_asset_returns(accept_licence=True, real=True)

    assert nominal.loc[pd.Period("2020", "Y"), "United States"] == pytest.approx(0.2)
    assert real.loc[pd.Period("2020", "Y"), "United States"] == pytest.approx(
        1.2 / 1.1 - 1, abs=1e-4
    )
    assert real.loc[pd.Period("2020", "Y"), "Germany"] == pytest.approx(0.0, abs=1e-4)


def test_millennium_data_names_second_columns_by_their_unit(monkeypatch):
    from financetoolkit import Economics
    from financetoolkit.economics import boe_model

    rows = [[None] * 5 for _ in range(7)]
    rows[3] = ["Description", "Bank Rate", None, "Share prices", "Consumer price index"]
    rows[5] = [
        "Units",
        "% end period",
        "%, calendar year average",
        "April 1962=100",
        "2015=100",
    ]
    rows += [[1720, 5, 5, 2.4, 0.4], [1721, 5, 5, 2.1, 0.41]]
    workbook = _workbook({"A1. Headline series": rows})
    monkeypatch.setattr(
        boe_model,
        "get_request",
        lambda url, timeout, extra_headers: FakeResponse(content=workbook),
    )

    millennium = Economics(
        start_date="1700-01-01", end_date="1800-12-31"
    ).get_millennium_of_macroeconomic_data()

    assert list(millennium.columns) == [
        "Bank Rate",
        "Bank Rate (%, calendar year average)",
        "Share prices",
        "Consumer price index",
    ]
    assert millennium.loc[pd.Period("1720", "Y"), "Bank Rate"] == pytest.approx(0.05)
    assert millennium.loc[pd.Period("1721", "Y"), "Share prices"] == pytest.approx(2.1)


def test_dnb_scenario_set_derives_zero_rates_from_the_state_variables(monkeypatch):
    import openpyxl

    from financetoolkit import Economics
    from financetoolkit.economics import dnb_model

    book = openpyxl.Workbook()
    book.remove(book.active)
    states = [[0.03, 0.04, 0.05], [0.03, 0.02, 0.01]]
    for number, sheet in enumerate(dnb_model.STATE_SHEETS):
        worksheet = book.create_sheet(sheet)
        for row in states:
            worksheet.append([value * (number + 1) for value in row])
    for sheet in dnb_model.RETURN_SHEETS.values():
        worksheet = book.create_sheet(sheet)
        worksheet.append([0.08, 0.06])
        worksheet.append([-0.02, 0.10])
    phi = book.create_sheet(dnb_model.NOMINAL_PHI)
    psi = book.create_sheet(dnb_model.NOMINAL_PSI)
    for maturity in range(1, 31):
        phi.append([-0.02 * maturity] * 3)
        psi.append([-0.1 * maturity, 0.0, 0.0])
    buffer = io.BytesIO()
    book.save(buffer)
    page = '<a href="/media/abc/cp2022-p-scenarioset-20k-2026q3.xlsx">P</a>'

    monkeypatch.setattr(
        dnb_model,
        "get_request",
        lambda url, timeout, extra_headers: (
            FakeResponse(content=buffer.getvalue())
            if url.endswith(".xlsx")
            else FakeResponse(text=page)
        ),
    )

    economics = Economics()
    scenarios = economics.get_scenario_set(
        variables=["Nominal rate 10Y", "Equity return"], scenarios=True, rounding=8
    )
    distribution = economics.get_scenario_set(
        variables="Equity return", quantiles=[0.5]
    )

    # exp(-(phi + psi x) / m) - 1 with phi = -0.2, psi = -1 and x = 0.04 in year 1.
    assert scenarios.loc[1, ("Nominal rate 10Y", 1)] == pytest.approx(
        np.exp(0.024) - 1, abs=1e-6
    )
    # Returns start in year 1.
    assert np.isnan(scenarios.loc[0, ("Equity return", 1)])
    assert scenarios.loc[2, ("Equity return", 2)] == pytest.approx(0.10)
    assert distribution.loc[1, ("Equity return", "50%")] == pytest.approx(0.03)
    assert economics.get_scenario_set(variables="Gold price").empty


def test_asset_class_proxies_are_period_returns(monkeypatch):
    from financetoolkit import Economics, historical_model

    days = pd.PeriodIndex(["2026-07-31", "2026-08-31", "2026-09-30"], freq="D")
    prices = pd.concat(
        {
            "Adj Close": pd.DataFrame(
                {"PSP": [50.0, 55.0, 49.5], "IGF": [60.0, 60.6, 60.0]}, index=days
            )
        },
        axis=1,
    )
    requested = []

    def fake(tickers, **kwargs):
        requested.append(tickers)
        return prices, []

    monkeypatch.setattr(historical_model, "get_historical_data", fake)

    proxies = Economics(
        start_date="2026-08-01", end_date="2026-09-30"
    ).get_asset_class_proxies(asset_classes=["Private Equity", "Infrastructure"])

    assert requested == [["PSP", "IGF"]]
    assert proxies.loc[pd.Period("2026-08", "M"), "Private Equity"] == pytest.approx(
        0.1
    )
    assert proxies.loc[pd.Period("2026-09", "M"), "Private Equity"] == pytest.approx(
        -0.1
    )
    assert Economics().get_asset_class_proxies(asset_classes="Art").empty


def test_eiopa_symmetric_adjustment_reads_the_daily_history(monkeypatch):
    from financetoolkit import FixedIncome
    from financetoolkit.fixedincome import eiopa_model

    rows = [[None] * 7 for _ in range(8)]
    rows += [
        [
            None,
            "All calendar days",
            None,
            "LT average",
            "(CI-AI) / AI",
            "Raw dampener",
            "Dampener final",
        ],
        [None, None, None, None, None, None, None],
        [None, pd.Timestamp("2026-09-30"), None, 1.30, 0.23, 0.077, 0.077],
        [None, pd.Timestamp("2026-09-29"), None, 1.30, 0.24, 0.079, 0.079],
        [None, pd.Timestamp("2020-03-16"), None, 1.20, -0.30, -0.19, -0.10],
    ]
    workbook = _workbook({"Calculations": rows})
    name = "EIOPA_symmetric_adjustment_equity_capital_charge"
    page = (
        f'<a href="/document/download/a_en?filename={name}_September_2026.xlsx">Sep</a>'
        f'<a href="/document/download/b_en?filename={name}_August_2026.xlsx">Aug</a>'
    )
    requested = []

    def fake(url, timeout):
        requested.append(url)
        return (
            FakeResponse(text=page)
            if url.endswith("_en")
            else FakeResponse(content=workbook)
        )

    monkeypatch.setattr(eiopa_model, "get_request", fake)

    adjustment = FixedIncome(
        start_date="2020-01-01", end_date="2026-09-30"
    ).get_eiopa_symmetric_adjustment()

    assert "September_2026" in requested[-1]
    assert adjustment.loc[
        pd.Period("2026-09-30", "D"), "Type 1 Equity Charge"
    ] == pytest.approx(0.467)
    assert adjustment.loc[
        pd.Period("2020-03-16", "D"), "Symmetric Adjustment"
    ] == pytest.approx(-0.10)
    assert adjustment.loc[
        pd.Period("2020-03-16", "D"), "Type 2 Equity Charge"
    ] == pytest.approx(0.39)


def test_option_chains_fall_back_to_cboe_when_yahoo_has_none(monkeypatch):
    import yfinance as yf

    from financetoolkit.economics import cboe_model
    from financetoolkit.options import options_model

    payload = {
        "data": {
            "current_price": 7818.93,
            "options": [
                {
                    "option": "SPX261016C07800000",
                    "bid": 80.0,
                    "ask": 82.0,
                    "iv": 0.12,
                    "open_interest": 900.0,
                },
                {
                    "option": "SPXW261016C07800000",
                    "bid": 80.5,
                    "ask": 82.5,
                    "iv": 0.121,
                    "open_interest": 50.0,
                },
                {
                    "option": "SPX261016C07900000",
                    "bid": 35.0,
                    "ask": 36.0,
                    "iv": 0.11,
                    "open_interest": 400.0,
                },
                {
                    "option": "SPX261016P07800000",
                    "bid": 60.0,
                    "ask": 61.0,
                    "iv": 0.13,
                    "open_interest": 700.0,
                },
                {
                    "option": "SPX261120C07800000",
                    "bid": 150.0,
                    "ask": 152.0,
                    "iv": 0.125,
                    "open_interest": 300.0,
                },
            ],
        }
    }

    class JsonResponse(FakeResponse):
        def json(self):
            return payload

    monkeypatch.setattr(cboe_model, "get_request", lambda url, timeout: JsonResponse())
    monkeypatch.setattr(
        yf, "Ticker", lambda ticker: (_ for _ in ()).throw(RuntimeError("down"))
    )

    assert options_model.get_option_expiry_dates("^SPX") == ["2026-10-16", "2026-11-20"]

    calls = options_model.get_option_chains(["^SPX"], "2026-10-16")

    # The monthly and the weekly contract on the same strike are one row: the more held.
    assert list(calls.loc["^SPX"].index) == [7800.0, 7900.0]
    assert calls.loc[("^SPX", 7800.0), "Contract Symbol"] == "SPX261016C07800000"
    assert calls.loc[("^SPX", 7800.0), "In The Money"]
    assert not calls.loc[("^SPX", 7900.0), "In The Money"]
    assert calls.loc[("^SPX", 7900.0), "Implied Volatility"] == pytest.approx(0.11)
